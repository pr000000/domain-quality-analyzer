"""
Website Domain Quality Analyzer (v0.1.2)
========================================
Author: Peiran (Cynthia) Jing | GitHub: https://github.com/pr000000

Overview
--------
Automatically analyzes the websites behind a list of customer accounts:
    - Reachability check (HTTPS/HTTP x www/non-www, with SSL fallback)
    - 18-dimension website quality score (0-24 points)
    - Industry keyword extraction (regex-based, English + Chinese)
    - Anomaly detection (bulk registration, duplicates, edu domains, mail providers)
    - Overseas-entity detection (language / ICP filing signals)

The results are appended as extra columns to the right of every sheet in the
original Excel file (or to the original CSV), so the input layout is preserved.

Data pipeline
-------------
    N raw records
        -> anomaly detection  (flag duplicates + bulk-registered domains)
        -> de-duplication     (unique domains only)
    Step 1   : reachability check (thread pool, batched)   -> reachable domains
    Step 1.5 : search-engine rescue for unreachable domains (extension point)
    Step 2   : 18-dimension scoring of reachable sites      -> score 0-24
    Step 3   : merge + write back to the original Excel/CSV format

Scoring system (18 dimensions, total 0-24)
------------------------------------------
    Layer 1 (HTML structure)   : 6 dimensions x 1 pt = 6
    Layer 2 (content richness) : 6 dimensions x 1 pt = 6
    Layer 3 (business signals) : 6 dimensions x 2 pt = 12
    Layer 3 is weighted double because it is the strongest indicator that a
    site belongs to a real, operating business rather than a parked page.
    Rating: ★★★ High (>=17) / ★★ Mid-High (>=10) / ★ Medium (>=5) / ☆ Low (0-4)

Usage
-----
    python domain_analysis.py <input_file_path>   (.xlsx or .csv)

Note: overseas domains are reached most reliably from a network node located
outside mainland China; other networks may produce more false "unreachable".
"""

import re
import csv
import json
import os
import sys
from collections import Counter, defaultdict
from typing import Optional


# ============================================================================
# Module 1: public mail detection
# ============================================================================

# Rule 2: brand keywords (with a trailing dot, so "outlook." matches
# outlook.com / outlook.jp but not e.g. "outlookgroup.com")
BRAND_KEYWORDS = [
    'outlook.', 'hotmail.', 'live.', 'msn.',
    'gmail.', 'googlemail.',
    'yahoo.', 'ymail.',
    'zohomail.', 'zoho.',
    'protonmail.', 'proton.',
    'icloud.', 'foxmail.', 'aliyun.',
    'aol.', 'yandex.', 'gmx.', 'tutanota.', 'fastmail.',
]

# Rule 3: exact-match set of common free mail domains (China + international)
SHORT_DOMAINS = {
    'qq.com', '163.com', '126.com', 'sina.com', 'sohu.com', '139.com',
    'me.com', 'mac.com', 'mail.ru', '189.cn', '2925.com',
}


def is_public_mail(domain: str) -> bool:
    """Return True if the domain belongs to a public / free mail provider.

    Such domains say nothing about the customer's own business website,
    so they are flagged instead of being scored.

    Rules (any one is sufficient):
        1. starts with 'mail.'
        2. contains any keyword in BRAND_KEYWORDS
        3. exact match in SHORT_DOMAINS

    Args:
        domain: e.g. 'qq.com', 'outlook.com', 'mail.example.com'
    """
    domain_lower = domain.lower().strip()
    # Rule 1
    if domain_lower.startswith('mail.'):
        return True
    # Rule 2
    for brand in BRAND_KEYWORDS:
        if brand in domain_lower:
            return True
    # Rule 3
    if domain_lower in SHORT_DOMAINS:
        return True
    return False


# ============================================================================
# Module 2: industry keyword patterns
# ============================================================================

# Each label maps to an English + Chinese regex, because the input list
# contains both domestic (Chinese) and international websites.
kw_patterns = {
    'AI': r'artificial.intelligence|machine.learning|深度学习|AI|GPT|LLM',
    'E-commerce': r'e-commerce|shop|store|购物|商城',
    'Gaming': r'game|gaming|游戏',
    'SaaS': r'saas|cloud.service|platform|平台',
    'Finance/Payments': r'fintech|payment|finance|金融|支付',
    'IT Services/Software': r'software|IT|技术|development',
    'Education': r'education|learning|course|教育|培训',
    'Healthcare': r'health|medical|pharma|医疗|健康',
    'Logistics': r'logistics|shipping|delivery|物流',
    'Media/Content': r'media|news|content|blog|媒体',
}


def extract_keywords(text: str, patterns: dict = None) -> str:
    """Extract matching industry labels from page text.

    Args:
        text: web page text / HTML
        patterns: label -> regex dict; defaults to kw_patterns

    Returns:
        Comma-separated labels, e.g. 'AI, SaaS' ('' if nothing matched)
    """
    if patterns is None:
        patterns = kw_patterns
    if not text:
        return ''
    matched = []
    for label, pattern in patterns.items():
        if re.search(pattern, text, re.IGNORECASE):
            matched.append(label)
    return ', '.join(matched)


# Mainland-China ICP / public-security filing numbers: a strong signal that
# the site is operated by a mainland Chinese entity (filing is mandatory there).
ICP_PATTERN = r'ICP备|ICP证|公网安备|京ICP|沪ICP|粤ICP|浙ICP|苏ICP|工信部'


def is_all_chinese_name(text: str) -> bool:
    """Return True if a company name is written entirely in Chinese.

    Definition: contains at least one CJK character and no Latin letters.
    Chinese/full-width punctuation, spaces and digits are allowed.
    Names containing Latin letters (e.g. "3M中国", "ABC科技") return False.
    """
    if not text:
        return False
    has_chinese = re.search(r'[\u4e00-\u9fff]', text) is not None
    has_latin = re.search(r'[A-Za-z]', text) is not None
    return has_chinese and not has_latin


def detect_overseas_entity(text: str, company: str = '') -> str:
    """Decide whether the entity behind a website is located overseas.

    Rules, from highest to lowest priority:
        0. company name is all Chinese         -> 'No'  (domestic; ignores page)
        1. no page content (unreachable/empty) -> ''    (cannot decide)
        2. ICP / public-security filing found  -> 'No'  (mainland China)
        3. <html lang> / meta language is zh   -> 'No'  (works for SPA sites)
        4. more kana/hangul than CJK chars     -> 'Yes' (Japanese/Korean site;
                                                  avoids counting kanji as Chinese)
        5. no Chinese in visible text, but Chinese elsewhere in the HTML
                                               -> 'No'  (JS-rendered Chinese site)
        6. no CJK characters anywhere          -> 'Yes' (overseas)

    Args:
        text: raw HTML (may be truncated)
        company: company / account name from the input file

    Returns:
        'Yes' / 'No' / '' (undetermined)
    """
    # Rule 0: an all-Chinese company name is treated as domestic. Highest
    # priority, and it also applies when the site is unreachable.
    if is_all_chinese_name(company):
        return 'No'
    # Rule 1
    if not text:
        return ''
    # Rule 2: strong signal - ICP filing => mainland China
    if re.search(ICP_PATTERN, text, re.IGNORECASE):
        return 'No'

    # Rule 3: language declarations. Even when the visible content is rendered
    # by JavaScript, the lang attribute still reveals the real language.
    # <html lang="zh-CN"> / <html lang="zh">
    if re.search(r'(?i)<html[^>]*\slang=["\']zh[\w-]*["\']', text):
        return 'No'
    # <meta http-equiv="content-language" content="zh-CN">
    if re.search(r'(?i)<meta[^>]*content-language[^>]*content=["\']zh[\w-]*["\']', text):
        return 'No'
    # <meta name="language" content="zh">  (non-standard but common)
    if re.search(r'(?i)<meta[^>]*name=["\']language["\'][^>]*content=["\']zh[\w-]*["\']', text):
        return 'No'

    # Rule 4: count CJK / Japanese / Korean characters in the full HTML
    chinese_all = re.findall(r'[\u4e00-\u9fff]', text)   # CJK ideographs
    kana = re.findall(r'[\u3040-\u30ff]', text)          # Japanese hiragana/katakana
    hangul = re.findall(r'[\uac00-\ud7a3]', text)        # Korean hangul syllables
    # More kana/hangul than ideographs => Japanese/Korean site (overseas);
    # its kanji/hanja should not be counted as Chinese.
    if (len(kana) + len(hangul)) > len(chinese_all):
        return 'Yes'

    # Rules 5/6: distinguish visible text from the full HTML
    visible = _visible_text(text)
    chinese_visible = re.findall(r'[\u4e00-\u9fff]', visible)

    # Chinese in visible text => domestic
    if len(chinese_visible) > 0:
        return 'No'

    # No Chinese in visible text, but Chinese inside scripts/JSON/comments
    # => SPA / front-end-rendered Chinese site
    if len(chinese_all) > 0:
        return 'No'

    # Rule 6: no CJK characters anywhere => overseas
    return 'Yes'


# ============================================================================
# Module 3: rating levels + Excel colour mapping
# ============================================================================

# Rating thresholds (minimum score for each level)
RATING_THRESHOLDS = {
    '★★★ High': 17,
    '★★ Mid-High': 10,
    '★ Medium': 5,
    '☆ Low': 0,
}

# Fill colour for each rating (openpyxl ARGB)
RATING_COLORS = {
    '★★★ High': 'FF92D050',     # green
    '★★ Mid-High': 'FF9BC2E6',  # blue
    '★ Medium': 'FFFFFF00',     # yellow
    '☆ Low': 'FFFF6B6B',        # red
}


def get_rating(score) -> Optional[str]:
    """Map a 0-24 score to its rating label (None if there is no score)."""
    if score is None:
        return None
    score = int(score)
    if score >= 17:
        return '★★★ High'
    elif score >= 10:
        return '★★ Mid-High'
    elif score >= 5:
        return '★ Medium'
    else:
        return '☆ Low'


def get_rating_fill(score):
    """Return the openpyxl PatternFill for a score's rating.

    Requires openpyxl; returns None if it is not installed (e.g. CSV output).
    """
    try:
        from openpyxl.styles import PatternFill
    except ImportError:
        return None

    rating = get_rating(score)
    if rating and rating in RATING_COLORS:
        color = RATING_COLORS[rating]
        return PatternFill(start_color=color, end_color=color, fill_type='solid')
    return None


# ============================================================================
# Module 4: anomaly detection (4 patterns)
# ============================================================================

# Educational / institutional domain suffixes
INSTITUTIONAL_PATTERNS = [r'\.edu(\.[a-z]{2,})?$', r'\.ac\.[a-z]{2,}$']


def detect_anomalies(records: list, domain_field: str = 'candidates',
                     bulk_tld_threshold: int = 3,
                     duplicate_threshold: int = 2) -> list:
    """Flag anomalous patterns in a list of domain records.

    Four anomaly types:
        1. Bulk registration - same base name under >= 3 different TLDs
        2. Institutional     - .edu.* / .ac.* suffix
        3. Mail provider     - public mail domain (see is_public_mail)
        4. Duplicate         - same domain used by >= 2 accounts
                               (the first occurrence is NOT flagged)

    Flagged records are excluded from reachability checks, which avoids
    wasting requests and keeps duplicate scores out of the statistics.

    Args:
        records: list of dicts; record[domain_field] is a list whose [0]
                 element is the primary domain
        domain_field: name of the domain field
        bulk_tld_threshold: TLD count that triggers "bulk registration"
        duplicate_threshold: occurrence count that triggers "duplicate"

    Returns:
        The same list, each record gaining an 'anomaly' field:
        None / 'Bulk registration' / 'Duplicate domain' /
        'Institutional domain' / 'Mail provider domain'
    """
    # Primary domain of every record
    domains = []
    for rec in records:
        candidates = rec.get(domain_field, [])
        domain = candidates[0] if candidates else ''
        domains.append(domain)

    # Anomaly 1: bulk registration - group TLDs by base name
    domain_groups = defaultdict(set)
    for d in domains:
        if not d:
            continue
        parts = d.rsplit('.', 1)
        if len(parts) == 2:
            base, tld = parts
            domain_groups[base].add(tld)
    batch_bases = {base for base, tlds in domain_groups.items()
                   if len(tlds) >= bulk_tld_threshold}

    # Anomaly 4: duplicates
    domain_counts = Counter(d for d in domains if d)
    seen_domains = set()

    # Label each record
    for i, rec in enumerate(records):
        domain = domains[i]
        if not domain:
            rec['anomaly'] = None
            continue

        base = domain.rsplit('.', 1)[0] if '.' in domain else domain

        # Priority: bulk registration > institutional > mail provider > duplicate
        if base in batch_bases:
            rec['anomaly'] = 'Bulk registration'
        elif any(re.search(p, domain) for p in INSTITUTIONAL_PATTERNS):
            rec['anomaly'] = 'Institutional domain'
        elif is_public_mail(domain):
            rec['anomaly'] = 'Mail provider domain'
        elif domain_counts[domain] >= duplicate_threshold and domain in seen_domains:
            rec['anomaly'] = 'Duplicate domain'
        else:
            rec['anomaly'] = None

        seen_domains.add(domain)

    return records


# ============================================================================
# Module 5: 18-dimension scoring (0-24)
# ============================================================================

# Browser version of the scorer: run inside a real browser (e.g. via
# Playwright's page.evaluate) so JavaScript-rendered content is included.
# Total 0-24: Layer 1 (6) + Layer 2 (6) + Layer 3 (12, 2 pts per dimension)
SCORING_JS = """(() => {
    let score = 0;
    const html = document.documentElement.innerHTML || '';
    const text = document.body ? document.body.innerText || '' : '';
    const title = document.title || '';

    // ===== Layer 1: HTML structure (max 6, 1 pt each) =====

    // 1.1 Valid title (not a 404 / error page)
    if (title && title.length > 3
        && !title.toLowerCase().includes('404')
        && !title.toLowerCase().includes('error')) score += 1;

    // 1.2 Non-trivial meta description
    const metaDesc = document.querySelector('meta[name="description"]');
    if (metaDesc && metaDesc.content && metaDesc.content.length > 10) score += 1;

    // 1.3 At least 5 links
    const links = document.querySelectorAll('a[href]');
    if (links.length >= 5) score += 1;

    // 1.4 Has images
    const images = document.querySelectorAll('img[src]');
    if (images.length >= 1) score += 1;

    // 1.5 Has navigation (nav or header)
    if (document.querySelector('nav') || document.querySelector('header')) score += 1;

    // 1.6 Has footer
    if (document.querySelector('footer')) score += 1;

    // ===== Layer 2: content richness (max 6, 1 pt each) =====

    // 2.1 Body text >= 200 characters
    if (text.length >= 200) score += 1;

    // 2.2 Has headings (h1-h3)
    const headings = document.querySelectorAll('h1, h2, h3');
    if (headings.length >= 1) score += 1;

    // 2.3 Has form elements
    const forms = document.querySelectorAll('form, input, button[type="submit"]');
    if (forms.length >= 1) score += 1;

    // 2.4 Has video / embedded content
    const media = document.querySelectorAll('video, iframe, embed');
    if (media.length >= 1) score += 1;

    // 2.5 Has list structures (ul/ol, at least 2)
    const lists = document.querySelectorAll('ul, ol');
    if (lists.length >= 2) score += 1;

    // 2.6 Has structured data (JSON-LD / microdata)
    const jsonLd = document.querySelectorAll('script[type="application/ld+json"]');
    const microdata = document.querySelectorAll('[itemscope]');
    if (jsonLd.length >= 1 || microdata.length >= 1) score += 1;

    // ===== Layer 3: business signals (max 12, 2 pts each) =====

    // 3.1 Contact information
    const contactPattern = /(?:tel:|phone|电话|联系|contact|email|@|地址|address)/i;
    if (contactPattern.test(html)) score += 2;

    // 3.2 Products / services / pricing
    const productPattern = /(?:product|service|solution|pricing|价格|产品|服务|方案)/i;
    if (productPattern.test(html)) score += 2;

    // 3.3 Social media links
    const socialPattern = /(?:facebook|twitter|linkedin|instagram|youtube|wechat|微信|weibo)/i;
    if (socialPattern.test(html)) score += 2;

    // 3.4 Call-to-action (sign up / trial / buy / contact)
    const ctaPattern = /(?:sign.?up|register|free.?trial|buy.?now|get.?started|contact.?us|注册|试用|立即|开始)/i;
    if (ctaPattern.test(html)) score += 2;

    // 3.5 Blog / news / resources
    const blogPattern = /(?:blog|news|resource|article|博客|新闻|资源|文章)/i;
    if (blogPattern.test(html)) score += 2;

    // 3.6 Privacy / terms pages (sign of a formally operated business)
    const legalPattern = /(?:privacy|terms|policy|legal|隐私|条款|协议)/i;
    if (legalPattern.test(html)) score += 2;

    return {
        score: score,
        title: title.substring(0, 100),
        text_len: text.length,
        link_count: links.length
    };
})()"""

# Human-readable description of every dimension
DIMENSIONS = {
    'Layer 1 - HTML structure (6 pts)': [
        '1.1 Valid title (not 404/error): 1 pt',
        '1.2 Meta description: 1 pt',
        '1.3 >= 5 links: 1 pt',
        '1.4 Has images: 1 pt',
        '1.5 Has navigation: 1 pt',
        '1.6 Has footer: 1 pt',
    ],
    'Layer 2 - Content richness (6 pts)': [
        '2.1 Body text >= 200 chars: 1 pt',
        '2.2 Has headings (h1-h3): 1 pt',
        '2.3 Has form elements: 1 pt',
        '2.4 Has video/embeds: 1 pt',
        '2.5 Has list structures: 1 pt',
        '2.6 Has structured data: 1 pt',
    ],
    'Layer 3 - Business signals (12 pts)': [
        '3.1 Contact information: 2 pts',
        '3.2 Products/services: 2 pts',
        '3.3 Social media: 2 pts',
        '3.4 Call-to-action: 2 pts',
        '3.5 Blog/resources: 2 pts',
        '3.6 Privacy/terms: 2 pts',
    ],
}

# Layer 3 regexes (kept identical to SCORING_JS so both scorers agree)
_L3_PATTERNS = {
    'contact': r'(?:tel:|phone|电话|联系|contact|email|@|地址|address)',
    'product': r'(?:product|service|solution|pricing|价格|产品|服务|方案)',
    'social': r'(?:facebook|twitter|linkedin|instagram|youtube|wechat|微信|weibo)',
    'cta': r'(?:sign.?up|register|free.?trial|buy.?now|get.?started|contact.?us|注册|试用|立即|开始)',
    'blog': r'(?:blog|news|resource|article|博客|新闻|资源|文章)',
    'legal': r'(?:privacy|terms|policy|legal|隐私|条款|协议)',
}


def _visible_text(html: str) -> str:
    """Roughly extract visible text (drop scripts, styles, tags, entities)."""
    h = re.sub(r'(?is)<(script|style|noscript)[^>]*>.*?</\1>', ' ', html)
    h = re.sub(r'(?s)<[^>]+>', ' ', h)
    h = re.sub(r'&[a-zA-Z#0-9]+;', ' ', h)
    return re.sub(r'\s+', ' ', h).strip()


def score_html(html: str) -> dict:
    """Score raw HTML on the 18 dimensions (0-24); Python port of SCORING_JS.

    Works on HTML fetched with requests, so no browser is needed and the
    whole list can be scored in parallel. Limitation: content that is only
    rendered client-side (SPA) is not visible here.

    Returns:
        {'score': int 0-24, 'title': str, 'text_len': int, 'link_count': int}
    """
    if not html:
        return {'score': 0, 'title': '', 'text_len': 0, 'link_count': 0}

    score = 0

    # -- Layer 1: HTML structure (1 pt each) --
    m = re.search(r'(?is)<title[^>]*>(.*?)</title>', html)
    title = re.sub(r'\s+', ' ', m.group(1).strip()) if m else ''
    tl = title.lower()
    if title and len(title) > 3 and '404' not in tl and 'error' not in tl:
        score += 1
    m = re.search(r'(?is)<meta[^>]+name=["\']description["\'][^>]*>', html)
    if m:
        cm = re.search(r'(?is)content=["\'](.*?)["\']', m.group(0))
        if cm and len(cm.group(1).strip()) > 10:
            score += 1
    link_count = len(re.findall(r'(?is)<a\s[^>]*href\s*=', html))
    if link_count >= 5:
        score += 1
    if len(re.findall(r'(?is)<img\s[^>]*src\s*=', html)) >= 1:
        score += 1
    if re.search(r'(?is)<nav[\s>]', html) or re.search(r'(?is)<header[\s>]', html):
        score += 1
    if re.search(r'(?is)<footer[\s>]', html):
        score += 1

    # -- Layer 2: content richness (1 pt each) --
    text = _visible_text(html)
    text_len = len(text)
    if text_len >= 200:
        score += 1
    if re.search(r'(?is)<h[1-3][\s>]', html):
        score += 1
    if (re.search(r'(?is)<form[\s>]', html) or re.search(r'(?is)<input[\s>]', html)
            or re.search(r'(?is)<button[^>]*type=["\']submit', html)):
        score += 1
    if re.search(r'(?is)<(video|iframe|embed)[\s>]', html):
        score += 1
    if len(re.findall(r'(?is)<(ul|ol)[\s>]', html)) >= 2:
        score += 1
    if re.search(r'(?is)application/ld\+json', html) or re.search(r'(?is)itemscope', html):
        score += 1

    # -- Layer 3: business signals (2 pts each) --
    for pat in _L3_PATTERNS.values():
        if re.search(pat, html, re.IGNORECASE):
            score += 2

    return {'score': score, 'title': title[:100],
            'text_len': text_len, 'link_count': link_count}


# ============================================================================
# Module 6: workflow steps
# ============================================================================

class DomainAnalysisWorkflow:
    """End-to-end domain analysis workflow.

    Step order: Step 0 -> Step 1 -> Step 1.5 -> Step 2 -> Step 3.
    Every step writes its intermediate results to JSON in output_dir, so a
    long run can be inspected or resumed step by step.

    Extension points:
    - Step 1   : ThreadPoolExecutor here; asyncio + aiohttp also works
    - Step 1.5 : plug in a search API (e.g. googlesearch-python, SerpAPI)
    - Step 2   : use Playwright + SCORING_JS for full browser rendering
    """

    def __init__(self, excel_path: str, output_dir: str = 'tmp'):
        """
        Args:
            excel_path: input file path (.xlsx or .csv)
            output_dir: directory for intermediate results
        """
        self.excel_path = excel_path
        self.output_dir = output_dir
        self.input_ext = os.path.splitext(excel_path)[1].lower()
        if self.input_ext not in ('.xlsx', '.csv'):
            raise ValueError(f"Unsupported file format: {self.input_ext} "
                             f"(only .xlsx and .csv are supported)")
        os.makedirs(output_dir, exist_ok=True)

    # ------------------------------------------------------------------
    # Step 0: preprocessing
    # ------------------------------------------------------------------
    def step0_preprocess(self) -> str:
        """Step 0: read the input, extract domain candidates, flag anomalies.

        Output: {output_dir}/batch_all.json with row_index, candidates and
        the anomaly flag of every record.

        Returns:
            path of the output JSON file
        """
        if self.input_ext == '.csv':
            all_records = self._read_csv_data()
        else:
            all_records = self._read_excel_data()

        # Run anomaly detection
        all_records = detect_anomalies(all_records, domain_field='candidates')

        output_path = os.path.join(self.output_dir, 'batch_all.json')
        with open(output_path, 'w', encoding='utf-8') as f:
            json.dump(all_records, f, ensure_ascii=False, indent=2)

        print(f"[Step 0] Preprocessing done: {len(all_records)} records -> {output_path}")
        return output_path

    def _read_csv_data(self) -> list:
        """Read records from a CSV file."""
        all_records = []

        # Try common encodings in order: CSVs exported on Chinese Windows
        # systems are often GBK rather than UTF-8.
        encodings_to_try = ['utf-8-sig', 'utf-8', 'gbk', 'gb2312', 'latin-1']
        rows = []
        headers = []
        for enc in encodings_to_try:
            try:
                with open(self.excel_path, 'r', encoding=enc, newline='') as f:
                    reader = csv.reader(f)
                    headers = next(reader)
                    rows = list(reader)
                break
            except (UnicodeDecodeError, UnicodeError):
                continue

        if not headers:
            raise ValueError(f"Could not read the CSV file (tried encodings: {encodings_to_try})")

        # Locate the domain column and company-name column
        domain_col_idx, company_col_idx = self._find_columns(headers)

        if domain_col_idx is None:
            raise ValueError(f"No domain column found in the CSV (header must contain "
                             f"domain/website/域名/网站). Current headers: {headers}")

        for row_idx, row in enumerate(rows, start=2):   # row 1 is the header
            raw_value = row[domain_col_idx] if domain_col_idx < len(row) else None
            if not raw_value:
                candidates = []
            else:
                candidates = self._extract_domain_candidates(str(raw_value))

            company = ''
            if company_col_idx is not None and company_col_idx < len(row):
                company = str(row[company_col_idx]) if row[company_col_idx] else ''

            all_records.append({
                'sheet': 'Sheet1',  # a CSV is treated as a single sheet
                'row_index': row_idx,
                'raw_value': str(raw_value) if raw_value else '',
                'candidates': candidates,
                'company': company,
            })

        return all_records

    def _read_excel_data(self) -> list:
        """Read records from every sheet of an Excel file."""
        import openpyxl

        wb = openpyxl.load_workbook(self.excel_path, data_only=True)
        all_records = []

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            headers = [cell.value for cell in ws[1]]

            domain_col_idx, company_col_idx = self._find_columns(headers)
            if domain_col_idx is None:
                continue   # skip sheets without a domain column

            for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                raw_value = row[domain_col_idx] if domain_col_idx < len(row) else None
                if not raw_value:
                    candidates = []
                else:
                    # Normalise to a bare domain (strip scheme, path, port)
                    candidates = self._extract_domain_candidates(str(raw_value))

                company = ''
                if company_col_idx is not None and company_col_idx < len(row):
                    company = str(row[company_col_idx]) if row[company_col_idx] else ''

                all_records.append({
                    'sheet': sheet_name,
                    'row_index': row_idx,
                    'raw_value': str(raw_value) if raw_value else '',
                    'candidates': candidates,
                    'company': company,
                })

        return all_records

    def _find_columns(self, headers: list) -> tuple:
        """Find the indices of the domain column and the company-name column.

        Matching is by keyword (English or Chinese) so the tool works with
        differently named exports.

        Returns:
            (domain_col_idx, company_col_idx); None where not found
        """
        # Domain column: header contains domain / website / 域名 / 网站
        domain_col_idx = None
        for i, h in enumerate(headers):
            if h and any(kw in str(h).lower() for kw in ['domain', 'website', '域名', '网站']):
                domain_col_idx = i
                break

        # Company column: header contains company / account / 公司 / 客户 / 账户
        company_col_idx = None
        for i, h in enumerate(headers):
            if h and any(kw in str(h).lower() for kw in ['company', 'account', '公司', '客户', '账户']):
                company_col_idx = i
                break
        # Fall back to column D (index 3), the usual position in our exports
        if company_col_idx is None and len(headers) > 3:
            company_col_idx = 3

        return domain_col_idx, company_col_idx

    def _extract_domain_candidates(self, raw: str) -> list:
        """Normalise a raw cell value into a list of domain candidates."""
        # Strip common scheme / www prefixes
        domain = raw.strip()
        for prefix in ['https://', 'http://', 'www.']:
            if domain.lower().startswith(prefix):
                domain = domain[len(prefix):]
        # Strip path
        domain = domain.split('/')[0]
        # Strip port
        domain = domain.split(':')[0]

        if not domain or '.' not in domain:
            return []
        return [domain.lower()]

    # ------------------------------------------------------------------
    # Step 1: reachability check
    # ------------------------------------------------------------------
    def step1_reachability_check(self, batch_json_path: str,
                                 batch_size: int = 80,
                                 max_workers: int = 15) -> list:
        """Step 1: check reachability of unique domains in parallel.

        Each domain is tried with up to 4 URL variants
        (https/http x bare/www); see check_url_variants.

        Args:
            batch_json_path: JSON produced by Step 0
            batch_size: domains per batch (each batch is saved separately)
            max_workers: number of parallel threads

        Returns:
            list of result file paths
        """
        with open(batch_json_path, 'r', encoding='utf-8') as f:
            records = json.load(f)

        # Unique domains, excluding records already flagged as anomalous
        unique_domains = list(set(
            rec['candidates'][0]
            for rec in records
            if rec['candidates'] and rec.get('anomaly') is None
        ))

        print(f"[Step 1] Unique domains: {len(unique_domains)}, "
              f"{(len(unique_domains) + batch_size - 1) // batch_size} batches x "
              f"{max_workers} workers")

        # Split into batches
        batches = [unique_domains[i:i + batch_size]
                   for i in range(0, len(unique_domains), batch_size)]

        from concurrent.futures import ThreadPoolExecutor

        result_files = []
        for batch_idx, batch_domains in enumerate(batches):
            output_file = os.path.join(self.output_dir, f'step1_results_{batch_idx}.json')
            result_files.append(output_file)

            # Network I/O bound -> threads give a near-linear speed-up
            with ThreadPoolExecutor(max_workers=max_workers) as executor:
                results = list(executor.map(self.check_url_variants, batch_domains))

            with open(output_file, 'w', encoding='utf-8') as f:
                json.dump(results, f, ensure_ascii=False, indent=2)

            reachable = sum(1 for r in results if r['accessible'])
            print(f"[Step 1] Batch {batch_idx}: {reachable}/{len(batch_domains)} reachable "
                  f"-> {output_file}")

        return result_files

    @staticmethod
    def _build_host_variants(domain: str) -> list:
        """Build www / non-www host candidates for a domain (ordered, de-duplicated).

        - example.com      -> [example.com, www.example.com]
        - www.example.com  -> [www.example.com, example.com]
        - shop.a.com       -> [shop.a.com, www.shop.a.com]

        This covers sites whose DNS record exists only on www or only on the
        bare domain, so neither case is misjudged as unreachable.
        """
        domain = domain.strip().lower().rstrip('.')
        if not domain:
            return []
        if domain.startswith('www.'):
            hosts = [domain, domain[4:]]
        else:
            hosts = [domain, 'www.' + domain]
        # Order-preserving de-duplication
        seen = set()
        return [h for h in hosts if h and not (h in seen or seen.add(h))]

    @staticmethod
    def check_url_variants(domain: str, timeout: int = 15) -> dict:
        """Check one domain across URL variants (www both ways + SSL fallback).

        Order (for a bare-domain input):
            1. https://{domain}        3. http://{domain}
            2. https://www.{domain}    4. http://www.{domain}
        If an https variant fails because of a certificate problem, it is
        retried once with verify=False, so a site that opens fine but has a
        mismatched certificate is not reported as unreachable.

        Args:
            domain: bare domain, e.g. 'example.com' or 'www.example.com'
            timeout: per-request timeout in seconds

        Returns:
            dict with domain, accessible, url_used (final URL after redirects),
            text (first 5000 chars), keywords, score fields, error, note
        """
        import requests
        try:
            import urllib3
            urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)
        except Exception:
            pass

        hosts = DomainAnalysisWorkflow._build_host_variants(domain)
        # All https hosts first, then all http hosts
        variants = [f'{scheme}{h}' for scheme in ('https://', 'http://') for h in hosts]

        headers = {
            'User-Agent': ('Mozilla/5.0 (Windows NT 10.0; Win64; x64) '
                           'AppleWebKit/537.36 (KHTML, like Gecko) '
                           'Chrome/124.0.0.0 Safari/537.36'),
            'Accept-Language': 'zh-CN,zh;q=0.9,en;q=0.8',
        }

        last_error = None
        for url in variants:
            is_https = url.startswith('https://')
            # https: verify the certificate first, then fall back to verify=False
            verify_attempts = (True, False) if is_https else (True,)
            for verify in verify_attempts:
                try:
                    resp = requests.get(url, timeout=timeout, allow_redirects=True,
                                        headers=headers, verify=verify)
                    # Encoding fix: requests defaults to ISO-8859-1 when no
                    # charset is declared, which garbles UTF-8 Chinese pages.
                    # Prefer the charset from the HTML <meta> tag.
                    if resp.encoding and resp.encoding.lower() in ('iso-8859-1', 'latin-1'):
                        meta_charset = re.search(
                            r'(?i)<meta[^>]*charset=["\']?([^"\'\s;>]+)',
                            resp.content[:2048].decode('ascii', errors='ignore'))
                        if meta_charset:
                            resp.encoding = meta_charset.group(1)
                        else:
                            resp.encoding = resp.apparent_encoding or 'utf-8'
                    body = resp.text or ''
                    if resp.status_code < 400 and len(body) > 100:
                        text = body[:5000]
                        sd = score_html(body)
                        return {
                            'domain': domain,
                            'accessible': True,
                            'url_used': resp.url or url,
                            'text': text,
                            'keywords': extract_keywords(text),
                            'score': sd['score'],
                            'title': sd['title'],
                            'text_len': sd['text_len'],
                            'link_count': sd['link_count'],
                            'error': None,
                            'note': '',
                        }
                    # Server responded but was blocked by a WAF / login wall /
                    # rate limiter => the site does exist, so count it as reachable
                    if resp.status_code in (401, 403, 429, 451, 503) and len(body) > 200:
                        text = body[:5000]
                        sd = score_html(body)
                        return {
                            'domain': domain,
                            'accessible': True,
                            'url_used': resp.url or url,
                            'text': text,
                            'keywords': extract_keywords(text),
                            'score': sd['score'],
                            'title': sd['title'],
                            'text_len': sd['text_len'],
                            'link_count': sd['link_count'],
                            'error': None,
                            'note': f'HTTP {resp.status_code}: likely WAF/login wall; '
                                    f'site exists, manual check recommended',
                        }
                    # Got a response but content is insufficient (404, empty
                    # page, ...): record it and move on to the next variant
                    last_error = f'HTTP {resp.status_code}, len={len(body)}'
                    break
                except requests.exceptions.SSLError:
                    last_error = 'SSL error'
                    continue  # retry with verify=False
                except Exception as e:
                    last_error = type(e).__name__
                    break  # connection/timeout error: try the next variant

        return {
            'domain': domain,
            'accessible': False,
            'url_used': None,
            'text': '',
            'keywords': '',
            'score': None,
            'title': '',
            'text_len': 0,
            'link_count': 0,
            'error': last_error or f'All {len(variants)} URL variants failed',
            'note': '',
        }

    # ------------------------------------------------------------------
    # Step 1.5: search-engine rescue of unreachable domains
    # ------------------------------------------------------------------
    def step1_5_rescue_unreachable(self, step1_results: list) -> str:
        """Step 1.5: try to rescue domains that Step 1 found unreachable.

        Idea: query a search engine with site:{domain} to confirm the site
        exists, and use the search snippets as supplementary keywords.

        Args:
            step1_results: list of Step 1 result file paths

        Returns:
            path of the rescue-result JSON file
        """
        unreachable = []
        for result_file in step1_results:
            if os.path.exists(result_file):
                with open(result_file, 'r', encoding='utf-8') as f:
                    results = json.load(f)
                unreachable.extend([r for r in results if not r['accessible']])

        print(f"[Step 1.5] Unreachable domains: {len(unreachable)}, starting search rescue...")

        rescued = []
        for item in unreachable:
            domain = item['domain']
            # ============================================================
            # TODO: implement the search rescue here, e.g.
            #   from googlesearch import search
            #   results = list(search(f"site:{domain}", num_results=3))
            #   if results:
            #       rescued.append({'domain': domain, 'rescued_url': results[0], ...})
            # ============================================================

        output_path = os.path.join(self.output_dir, 'step1_5_rescued.json')
        with open(output_path, 'w', encoding='utf-8') as f:
            json.dump(rescued, f, ensure_ascii=False, indent=2)

        print(f"[Step 1.5] Rescue done: {len(rescued)} domains rescued")
        return output_path

    # ------------------------------------------------------------------
    # Step 2: 18-dimension scoring
    # ------------------------------------------------------------------
    def step2_browser_scoring(self, accessible_urls: list,
                              batch_size: int = 15) -> list:
        """Step 2: 18-dimension scoring (0-24).

        Reuses the score already computed by score_html during Step 1, so no
        second request or browser launch is needed. Items without a score
        (e.g. from the rescue step) are scored from their text on the spot.

        Args:
            accessible_urls: reachable items (Step 1 result dicts)
            batch_size: items per saved batch file

        Returns:
            list of scoring result file paths
        """
        print(f"[Step 2] Scoring {len(accessible_urls)} reachable sites, "
              f"{batch_size} per batch")

        result_files = []
        for batch_idx in range(0, len(accessible_urls), batch_size):
            batch = accessible_urls[batch_idx:batch_idx + batch_size]
            batch_results = []

            for item in batch:
                url = item.get('url_used')
                domain = item['domain']

                # Prefer the Step 1 score; otherwise score the stored text now
                score = item.get('score')
                if score is None and item.get('text'):
                    sd = score_html(item['text'])
                    score = sd['score']
                    title = sd['title']
                    text_len = sd['text_len']
                    link_count = sd['link_count']
                else:
                    title = item.get('title', '')
                    text_len = item.get('text_len', 0)
                    link_count = item.get('link_count', 0)

                batch_results.append({
                    'domain': domain,
                    'url': url,
                    'score': score,
                    'title': title,
                    'text_len': text_len,
                    'link_count': link_count,
                    'rating': get_rating(score),
                })

            output_file = os.path.join(self.output_dir, f'step2_batch_{batch_idx // batch_size}.json')
            with open(output_file, 'w', encoding='utf-8') as f:
                json.dump(batch_results, f, ensure_ascii=False, indent=2)
            result_files.append(output_file)

        return result_files

    # ------------------------------------------------------------------
    # Step 3: merge and output
    # ------------------------------------------------------------------
    # Columns appended to the right of the original data
    NEW_COLUMNS = ['Reachable', 'URL Used', 'Score (0-24)', 'Rating',
                   'Industry Keywords', 'Overseas Entity', 'Anomaly', 'Note']

    def step3_merge_and_output(self) -> str:
        """Step 3: merge all intermediate results into the original file format.

        - Excel input -> columns appended to every sheet, saved as *_analyzed.xlsx
        - CSV input   -> columns appended to the rows, saved as *_analyzed.csv

        Returns:
            output file path
        """
        from glob import glob

        # Load Step 0 records
        batch_all_path = os.path.join(self.output_dir, 'batch_all.json')
        with open(batch_all_path, 'r', encoding='utf-8') as f:
            all_records = json.load(f)

        # Step 1 results: domain -> {accessible, url_used, keywords, ...}
        reachability_map = {}
        for fpath in sorted(glob(os.path.join(self.output_dir, 'step1_results_*.json'))):
            with open(fpath, 'r', encoding='utf-8') as f:
                for item in json.load(f):
                    reachability_map[item['domain']] = item

        # Step 1.5 rescue results override Step 1 for the same domain
        rescue_path = os.path.join(self.output_dir, 'step1_5_rescued.json')
        if os.path.exists(rescue_path):
            with open(rescue_path, 'r', encoding='utf-8') as f:
                for item in json.load(f):
                    reachability_map[item['domain']] = item

        # Step 2 results: domain -> {score, rating, title, ...}
        scoring_map = {}
        for fpath in sorted(glob(os.path.join(self.output_dir, 'step2_batch_*.json'))):
            with open(fpath, 'r', encoding='utf-8') as f:
                for item in json.load(f):
                    scoring_map[item['domain']] = item

        if self.input_ext == '.csv':
            return self._output_csv(all_records, reachability_map, scoring_map)
        else:
            return self._output_excel(all_records, reachability_map, scoring_map)

    def _analysis_row(self, rec, reachability_map, scoring_map) -> list:
        """Build the appended column values for one record."""
        domain = rec['candidates'][0] if rec['candidates'] else ''
        reach_info = reachability_map.get(domain, {})
        score_info = scoring_map.get(domain, {})

        accessible = 'Yes' if reach_info.get('accessible') else 'No'
        url_used = reach_info.get('url_used', '') or ''
        score = score_info.get('score')
        rating = get_rating(score)
        keywords = reach_info.get('keywords', '') or score_info.get('keywords', '')
        overseas = detect_overseas_entity(reach_info.get('text', ''),
                                          rec.get('company', ''))
        anomaly = rec.get('anomaly') or ''
        note = reach_info.get('note', '')
        return [accessible, url_used, score, rating, keywords, overseas, anomaly, note]

    def _output_excel(self, all_records, reachability_map, scoring_map) -> str:
        """Append analysis columns to every sheet of the original Excel file."""
        import openpyxl
        from openpyxl.styles import Font

        wb = openpyxl.load_workbook(self.excel_path)

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            max_col = ws.max_column

            # Header row
            for col_offset, col_name in enumerate(self.NEW_COLUMNS, start=1):
                cell = ws.cell(row=1, column=max_col + col_offset, value=col_name)
                cell.font = Font(bold=True)

            # Data rows
            sheet_records = [r for r in all_records if r['sheet'] == sheet_name]
            for rec in sheet_records:
                values = self._analysis_row(rec, reachability_map, scoring_map)
                score, rating = values[2], values[3]
                for col_offset, val in enumerate(values, start=1):
                    cell = ws.cell(row=rec['row_index'], column=max_col + col_offset, value=val)
                    # Colour-code the rating column
                    if col_offset == 4 and rating:
                        fill = get_rating_fill(score)
                        if fill:
                            cell.fill = fill

        output_path = self.excel_path.replace('.xlsx', '_analyzed.xlsx')
        wb.save(output_path)
        print(f"[Step 3] Output written: {output_path}")
        return output_path

    def _output_csv(self, all_records, reachability_map, scoring_map) -> str:
        """Append analysis columns to the original CSV and write a new CSV."""
        # Re-read the original CSV (same encoding fallback as Step 0)
        encodings_to_try = ['utf-8-sig', 'utf-8', 'gbk', 'gb2312', 'latin-1']
        original_rows = []
        original_headers = []
        for enc in encodings_to_try:
            try:
                with open(self.excel_path, 'r', encoding=enc, newline='') as f:
                    reader = csv.reader(f)
                    original_headers = next(reader)
                    original_rows = list(reader)
                break
            except (UnicodeDecodeError, UnicodeError):
                continue

        # row_index -> appended values
        result_map = {}
        for rec in all_records:
            values = self._analysis_row(rec, reachability_map, scoring_map)
            if values[2] is None:
                values[2] = ''
            result_map[rec['row_index']] = values

        output_path = self.excel_path.replace('.csv', '_analyzed.csv')
        # utf-8-sig so Excel opens the Chinese/star characters correctly
        with open(output_path, 'w', encoding='utf-8-sig', newline='') as f:
            writer = csv.writer(f)
            writer.writerow(original_headers + self.NEW_COLUMNS)
            for i, row in enumerate(original_rows, start=2):
                analysis = result_map.get(i, [''] * len(self.NEW_COLUMNS))
                writer.writerow(row + analysis)

        print(f"[Step 3] CSV output written: {output_path}")
        return output_path


# ============================================================================
# CLI entry point
# ============================================================================

def main():
    """Command-line entry point."""
    print("=" * 60)
    print("  Website Domain Quality Analyzer v0.1.2")
    print("=" * 60)

    if len(sys.argv) < 2:
        print("Usage: python domain_analysis.py <input_file_path>")
        print()
        print("  input_file_path: file containing customer domains (.xlsx or .csv)")
        print()
        print("Examples:")
        print("  python domain_analysis.py ./data/accounts.xlsx")
        print("  python domain_analysis.py ./data/accounts.csv")
        sys.exit(1)

    input_path = sys.argv[1]
    if not os.path.exists(input_path):
        print(f"Error: file not found - {input_path}")
        sys.exit(1)

    ext = os.path.splitext(input_path)[1].lower()
    if ext not in ('.xlsx', '.csv'):
        print(f"Error: unsupported format '{ext}' (only .xlsx and .csv)")
        sys.exit(1)

    workflow = DomainAnalysisWorkflow(input_path)

    print("\n" + "-" * 40 + "\nStep 0: Preprocessing\n" + "-" * 40)
    batch_json = workflow.step0_preprocess()

    print("\n" + "-" * 40 + "\nStep 1: Reachability check (parallel)\n" + "-" * 40)
    step1_files = workflow.step1_reachability_check(batch_json)

    print("\n" + "-" * 40 + "\nStep 1.5: Search rescue of unreachable domains\n" + "-" * 40)
    rescue_file = workflow.step1_5_rescue_unreachable(step1_files)

    print("\n" + "-" * 40 + "\nStep 2: 18-dimension scoring\n" + "-" * 40)
    # Collect reachable items from Step 1 + Step 1.5
    accessible_urls = []
    for fpath in step1_files:
        if os.path.exists(fpath):
            with open(fpath, 'r', encoding='utf-8') as f:
                accessible_urls.extend([r for r in json.load(f) if r['accessible']])
    if os.path.exists(rescue_file):
        with open(rescue_file, 'r', encoding='utf-8') as f:
            accessible_urls.extend(json.load(f))

    workflow.step2_browser_scoring(accessible_urls)

    print("\n" + "-" * 40 + "\nStep 3: Merge and output\n" + "-" * 40)
    output_file = workflow.step3_merge_and_output()

    print()
    print("=" * 60)
    print(f"  Done! Output file: {output_file}")
    print("=" * 60)


if __name__ == '__main__':
    main()
